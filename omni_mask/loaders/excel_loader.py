import openpyxl
from typing import Any, Set

from omni_mask.loaders.base import BaseLoader


class ExcelLoader(BaseLoader):
    def can_handle(self, filepath: str) -> bool:
        return filepath.lower().endswith((".xlsx", ".xls"))

    def anonymize(
        self,
        filepath: str,
        outpath: str,
        core: Any,
        pii_enabled: Set = None,
        enabled_fastmask: Set = None,
    ) -> None:
        core.reset_records()
        wb = openpyxl.load_workbook(filepath)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        cell.value = _process_segment(
                            cell.value, core, pii_enabled, enabled_fastmask
                        )
        wb.save(outpath)
        wb.close()

    def deanonymize(self, filepath: str, outpath: str, core: Any) -> None:
        wb = openpyxl.load_workbook(filepath)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        cell.value = core.deanonymize_text(cell.value)
        wb.save(outpath)
        wb.close()


def _process_segment(
    text: str, core: Any, pii_enabled: Set, enabled_fastmask: Set
) -> str:
    if pii_enabled:
        text, pii_mappings = core.pii_anonymize_text(text, pii_enabled)
        core.accumulate_pii_mappings(pii_mappings)
    if enabled_fastmask:
        rules = core._build_fastmask_rules(enabled_fastmask)
        if rules:
            fm_masker = __import__(
                "llm_router_plugins.maskers.fast_masker.core.masker",
                fromlist=["FastMasker"],
            ).FastMasker(rules)
            text, fm_mappings = fm_masker.mask(text)
            core.accumulate_fastmask_mappings(fm_mappings)
    return text
