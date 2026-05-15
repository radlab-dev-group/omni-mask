from typing import Any, Set
from docx import Document

from omni_mask.loaders.base import BaseLoader


class DocxLoader(BaseLoader):
    def can_handle(self, filepath: str) -> bool:
        return filepath.lower().endswith((".docx", ".doc"))

    def anonymize(
        self,
        filepath: str,
        outpath: str,
        core: Any,
        pii_enabled: Set = None,
        enabled_fastmask: Set = None,
    ) -> None:
        core.reset_records()
        doc = Document(filepath)

        for para in doc.paragraphs:
            for run in para.runs:
                if run.text:
                    run.text = _process_segment(
                        run.text, core, pii_enabled, enabled_fastmask
                    )

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if run.text:
                                run.text = _process_segment(
                                    run.text, core, pii_enabled, enabled_fastmask
                                )

        doc.save(outpath)

    def deanonymize(self, filepath: str, outpath: str, core: Any) -> None:
        doc = Document(filepath)
        for para in doc.paragraphs:
            for run in para.runs:
                if run.text:
                    run.text = core.deanonymize_text(run.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if run.text:
                                run.text = core.deanonymize_text(run.text)
        doc.save(outpath)


def _process_segment(
    text: str, core: Any, pii_enabled: Set, enabled_fastmask: Set
) -> str:
    if pii_enabled:
        text, pii_mappings = core.pii_anonymize_text(text, pii_enabled)
        core.accumulate_pii_mappings(pii_mappings)
    if enabled_fastmask:
        rules = core._build_fastmask_rules(enabled_fastmask)
        if rules:
            fm_masker = __import__(
                "llm_router_plugins.maskers.fast_masker.core.masker",
                fromlist=["FastMasker"],
            ).FastMasker(rules)
            text, fm_mappings = fm_masker.mask(text)
            core.accumulate_fastmask_mappings(fm_mappings)
    return text
